from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pt_invoice_adder.list_writer import append_rows


def _make_list(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "工作表1"
    ws.append(
        [
            "Invoice Date",
            "Country",
            "Shipment type",
            "GMBH INV.",
            "total (總件數)",
            "gross weight(KG)",
        ]
    )
    ws.append([dt.datetime(2026, 1, 1), "DE", "SEA", 111, 1, 10.0])
    wb.save(path)


def test_append_and_skip_duplicate(tmp_path: Path):
    xlsx = tmp_path / "PT INV LIST.xlsx"
    _make_list(xlsx)

    rows = [
        {
            "invoice_date": dt.date(2026, 7, 8),
            "country": "MY",
            "shipment_type": "SEA",
            "invoice_no": "50656407",
            "total_pkg": 5,
            "gross_weight": 437.53,
        },
        {
            "invoice_date": dt.date(2026, 1, 1),
            "country": "DE",
            "shipment_type": "SEA",
            "invoice_no": "111",  # duplicate
            "total_pkg": 9,
            "gross_weight": 99.0,
        },
    ]
    added, skipped = append_rows(xlsx, rows)
    assert added == 1
    assert skipped == 1

    wb = load_workbook(xlsx)
    ws = wb["工作表1"]
    assert ws.max_row == 3
    assert ws.cell(3, 4).value == 50656407
    assert ws.cell(3, 2).value == "MY"
    assert ws.cell(3, 5).value == 5
    assert float(ws.cell(3, 6).value) == 437.53

    added2, skipped2 = append_rows(xlsx, rows, dry_run=True)
    assert added2 == 0
    assert skipped2 == 2
