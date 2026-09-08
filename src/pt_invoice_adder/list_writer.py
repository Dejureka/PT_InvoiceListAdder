"""Append extracted invoice rows to PT INV LIST.xlsx (columns A–F)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


HEADERS = (
    "Invoice Date",
    "Country",
    "Shipment type",
    "GMBH INV.",
    "total (總件數)",
    "gross weight(KG)",
)


def _get_sheet(wb) -> Worksheet:
    if "工作表1" in wb.sheetnames:
        return wb["工作表1"]
    return wb[wb.sheetnames[0]]


def _norm_inv(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    s = str(value).strip()
    if s.endswith(".0"):
        try:
            return str(int(float(s)))
        except ValueError:
            pass
    return s


def existing_inv_set(ws: Worksheet) -> set[str]:
    found: set[str] = set()
    for r in range(2, (ws.max_row or 1) + 1):
        inv = _norm_inv(ws.cell(r, 4).value)
        if inv:
            found.add(inv)
    return found


def append_rows(
    list_path: str | Path,
    rows: Sequence[dict[str, Any]],
    *,
    dry_run: bool = False,
) -> tuple[int, int]:
    """Append A–F only; skip duplicate GMBH INV in col D. Return (added, skipped)."""
    list_path = Path(list_path)
    if list_path.exists():
        wb = load_workbook(list_path)
        ws = _get_sheet(wb)
    else:
        if dry_run:
            return (len(rows), 0)
        wb = Workbook()
        ws = wb.active
        ws.title = "工作表1"
        for c, h in enumerate(HEADERS, 1):
            ws.cell(1, c, h)

    existing = existing_inv_set(ws)
    added = 0
    skipped = 0
    next_row = (ws.max_row or 1) + 1
    # if sheet only has empty first row without headers detection — still fine
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        for c, h in enumerate(HEADERS, 1):
            ws.cell(1, c, h)
        next_row = 2

    for row in rows:
        inv = _norm_inv(row.get("invoice_no"))
        if not inv or inv in existing:
            skipped += 1
            continue
        if dry_run:
            added += 1
            existing.add(inv)
            continue

        inv_date = row.get("invoice_date")
        if isinstance(inv_date, date) and not isinstance(inv_date, datetime):
            cell_date: Any = datetime(inv_date.year, inv_date.month, inv_date.day)
        else:
            cell_date = inv_date

        inv_cell: Any = inv
        try:
            inv_cell = int(inv)
        except ValueError:
            pass

        ws.cell(next_row, 1, cell_date)
        ws.cell(next_row, 2, row.get("country"))
        ws.cell(next_row, 3, row.get("shipment_type"))
        ws.cell(next_row, 4, inv_cell)
        ws.cell(next_row, 5, row.get("total_pkg"))
        ws.cell(next_row, 6, row.get("gross_weight"))
        existing.add(inv)
        next_row += 1
        added += 1

    if not dry_run and added:
        wb.save(list_path)
    return added, skipped
